from openpyxl import Workbook, load_workbook
import os

FILE_NAME = "student_fees.xlsx"

def add_student():
    roll = input("Enter Roll No: ")
    name = input("Enter Student Name: ")
    course = input("Enter Course: ")
    fees = input("Enter Fees: ")

    if os.path.exists(FILE_NAME):
        wb = load_workbook(FILE_NAME)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["Roll No", "Name", "Course", "Fees"])

    ws.append([roll, name, course, fees])
    wb.save(FILE_NAME)

    print("Student fees saved successfully!")

def show_students():
    if not os.path.exists(FILE_NAME):
        print("No student data found.")
        return

    wb = load_workbook(FILE_NAME)
    ws = wb.active

    for row in ws.iter_rows(values_only=True):
        print(row)

while True:
    print("\n--- Student Fees Management System ---")
    print("1. Add Student")
    print("2. Show Students")
    print("3. Exit")

    choice = input("Enter choice: ")

    if choice == "1":
        add_student()
    elif choice == "2":
        show_students()
    elif choice == "3":
        print("Thank you!")
        break
    else:
        print("Invalid choice")